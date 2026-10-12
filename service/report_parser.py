"""报告文件解析：文本抽取 + 指标识别。

支持的上传格式：
* **PDF**    —— pdfminer.six 优先，失败回退 pypdf
* **文本**   —— txt / csv / md（UTF-8，回退 GBK）
* **图片**   —— png / jpg / webp（本环境无 OCR，仅保存附件并提示手工补录）
* **Word/Excel** —— 若安装了 python-docx / openpyxl 则解析，否则给出提示

指标识别采用「正则为主 + 大模型兜底」：
先按行解析「名称 结果 单位 参考范围 异常标记」的常见排版，
解析不到时再由 ``ReportService`` 调用大模型抽取。
"""
from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}

# ----------------------------------------------------------------------
# 报告名称 / 类型识别
# ----------------------------------------------------------------------
LAB_KEYWORDS = [
    "血常规", "尿常规", "便常规", "肝肾功能", "肝功能", "肾功能", "血脂", "血糖",
    "电解质", "心肌酶", "甲状腺功能", "甲功", "凝血", "C反应蛋白", "糖化血红蛋白",
    "血沉", "乙肝", "肿瘤标志物", "血气分析",
]
EXAM_KEYWORDS = [
    "胸部CT", "CT", "核磁", "MRI", "X线", "DR", "超声", "彩超", "B超",
    "心电图", "胃镜", "肠镜", "胸片", "造影", "PET",
]

# 常见单位（长单位放前面，避免被短单位抢先匹配）
UNITS = [
    "10^12/L", "10^9/L", "10^6/L", "×10^12/L", "×10^9/L", "×10^6/L",
    "mmol/L", "μmol/L", "umol/L", "mg/dL", "mg/L", "ng/mL", "pg/mL", "g/L",
    "U/L", "IU/L", "mIU/L", "mmHg", "次/分", "个/HP", "个/hp", "×10^9",
    "fL", "pg", "ng", "%", "‰",
]

# 全角 → 半角
_FULL_HALF = {
    "（": "(", "）": ")", "：": ":", "％": "%", "～": "~", "－": "-", "＞": ">", "＜": "<",
    "，": ",", "、": ",", "。": ".", "　": " ", "／": "/", "＋": "+", "℃": "C",
}
_FLAG_WORDS = ("偏高", "偏低", "升高", "降低", "正常", "异常", "阳性", "阴性", "弱阳性")

_RANGE_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:~|-|—|–|至|--)\s*(\d+(?:\.\d+)?)")
_NUM_RE = re.compile(r"(?<![\d.])(\d+(?:\.\d+)?)(?![\d.])")
_CODE_RE = re.compile(r"[（(]?\b([A-Z][A-Z0-9%\-]{1,9})\b[)）]?")
_HEADER_WORDS = ("项目", "名称", "结果", "单位", "参考范围", "参考值", "提示", "序号", "检测")


def normalize(text: str) -> str:
    for src, dst in _FULL_HALF.items():
        text = text.replace(src, dst)
    return re.sub(r"[ \t\u00a0]+", " ", text).strip()


# ----------------------------------------------------------------------
# 文本抽取
# ----------------------------------------------------------------------
def extract_text(filename: str, data: bytes) -> Tuple[str, str]:
    """抽取纯文本。

    :return: (文本, 说明信息)
    """
    ext = Path(filename).suffix.lower()

    if ext in IMAGE_EXT:
        return "", "图片报告已保存，但当前环境未启用 OCR，请手工补录指标或改用 PDF/文本格式"

    if ext == ".pdf":
        return _extract_pdf(data)

    if ext in (".txt", ".csv", ".md", ".tsv", ".json"):
        return _decode_bytes(data), "已按文本解析"

    if ext == ".docx":
        return _extract_docx(data)

    if ext in (".xlsx", ".xls"):
        return _extract_xlsx(data)

    return "", f"暂不支持解析 {ext or '该'} 格式，已保存附件"


def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", "replace")


def _extract_pdf(data: bytes) -> Tuple[str, str]:
    try:
        from pdfminer.high_level import extract_text as pdf_extract

        text = pdf_extract(io.BytesIO(data)) or ""
        if text.strip():
            return text, "已通过 pdfminer 解析 PDF"
    except Exception as exc:  # noqa: BLE001
        logger.info("pdfminer 解析失败，改用 pypdf：%s", exc)

    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages[:30]]
        text = "\n".join(pages)
        if text.strip():
            return text, "已通过 pypdf 解析 PDF"
        return "", "PDF 中未提取到文本层（可能是扫描件），请手工补录指标"
    except Exception as exc:  # noqa: BLE001
        logger.warning("PDF 解析失败：%s", exc)
        return "", f"PDF 解析失败：{exc}"


def _extract_docx(data: bytes) -> Tuple[str, str]:
    try:
        import docx  # type: ignore

        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" ".join(cell.text for cell in row.cells))
        text = "\n".join(lines)
        return (text, "已解析 Word 文档") if text.strip() else ("", "Word 文档中未提取到文本")
    except ImportError:
        return "", "未安装 python-docx，无法解析 Word 文档，请在服务器安装后重试"
    except Exception as exc:  # noqa: BLE001
        return "", f"Word 解析失败：{exc}"


def _extract_xlsx(data: bytes) -> Tuple[str, str]:
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        lines: List[str] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" ".join(cells))
        text = "\n".join(lines)
        return (text, "已解析 Excel 表格") if text.strip() else ("", "Excel 中未提取到数据")
    except ImportError:
        return "", "未安装 openpyxl，无法解析 Excel，请在服务器安装后重试"
    except Exception as exc:  # noqa: BLE001
        return "", f"Excel 解析失败：{exc}"


# ----------------------------------------------------------------------
# 指标识别
# ----------------------------------------------------------------------
def parse_items(text: str) -> List[Dict[str, Any]]:
    """从报告文本中解析指标明细。"""
    items: List[Dict[str, Any]] = []
    seen: set = set()

    for raw in (text or "").splitlines():
        line = normalize(raw)
        if len(line) < 4 or len(line) > 200:
            continue
        if _looks_like_header(line):
            continue
        item = _parse_line(line)
        if not item:
            continue
        key = (item["name"], item["value"])
        if key in seen:
            continue
        seen.add(key)
        item["sort_order"] = len(items)
        items.append(item)
        if len(items) >= 200:
            break

    return items


def _looks_like_header(line: str) -> bool:
    if _NUM_RE.search(line):
        return False
    return sum(1 for w in _HEADER_WORDS if w in line) >= 2


def _parse_line(line: str) -> Optional[Dict[str, Any]]:
    """解析单行指标。返回 None 表示该行不像指标。"""
    explicit_flag = None
    for word in _FLAG_WORDS:
        if word in line:
            explicit_flag = word
            break
    if "↑" in line:
        explicit_flag = explicit_flag or "偏高"
    if "↓" in line:
        explicit_flag = explicit_flag or "偏低"

    working = line.replace("↑", " ").replace("↓", " ")

    # 1) 参考范围
    low = high = None
    ref_text = None
    m_range = _RANGE_RE.search(working)
    if m_range:
        try:
            low, high = float(m_range.group(1)), float(m_range.group(2))
            ref_text = f"{_trim(m_range.group(1))}-{_trim(m_range.group(2))}"
        except ValueError:
            low = high = None
        working = working[: m_range.start()] + " " + working[m_range.end():]

    # 2) 结果值 + 单位
    value, unit, value_span = _find_value_and_unit(working)
    if value is None or value_span is None:
        return None

    # 3) 名称 = 结果值之前的部分
    name = normalize(working[: value_span[0]])
    name = re.sub(r"[:：\-–—,，、\s]+$", "", name).strip(" :,-")
    name = re.sub(r"^[\d.\s]+", "", name)
    if not name or len(name) < 2:
        return None
    if _NUM_RE.fullmatch(name):
        return None

    # 4) 指标编码
    code = None
    m_code = _CODE_RE.search(name)
    if m_code:
        code = m_code.group(1)

    name = re.sub(r"\s+", " ", name)[:100]

    # 5) 异常标记
    flag = _resolve_flag(value, low, high, explicit_flag)

    # 至少要有一个「范围」或「单位」才认为是指标行，避免噪声
    if ref_text is None and unit is None:
        return None

    return {
        "name": name,
        "code": code,
        "value": value,
        "unit": unit,
        "ref": ref_text,
        "flag": flag,
    }


_UNIT_RE = re.compile("(?:×)?" + "(" + "|".join(re.escape(u) for u in UNITS) + ")")


def _find_value_and_unit(working: str) -> Tuple[Optional[str], Optional[str], Optional[tuple]]:
    """在一行中定位「结果值」与紧随其后的「单位」。

    关键点：单位可能出现在名称里（如 ``中性粒细胞百分比 NEUT%``），
    因此先把所有单位片段的位置算出来，再排除落在单位内部的数字，
    避免把 ``10^9/L`` 里的 ``10`` 当成结果值。
    """
    unit_spans = [(m.start(), m.end(), m.group(1)) for m in _UNIT_RE.finditer(working)]
    numbers = [
        m for m in _NUM_RE.finditer(working)
        if not any(start <= m.start() < end for start, end, _ in unit_spans)
    ]
    if not numbers:
        return None, None, None

    # 优先：数字后面紧跟单位
    for m in numbers:
        tail = working[m.end(): m.end() + 16].lstrip()
        m_unit = _UNIT_RE.match(tail)
        if m_unit:
            return m.group(1), m_unit.group(1), (m.start(), m.end())

    # 其次：行尾最后一个数字（无单位，但通常有参考范围）
    last = numbers[-1]
    if not working[last.end():].strip():
        return last.group(1), None, (last.start(), last.end())

    # 兜底：取第一个数字
    first = numbers[0]
    return first.group(1), None, (first.start(), first.end())


def _resolve_flag(value: str, low: Optional[float], high: Optional[float],
                  explicit: Optional[str]) -> str:
    if explicit in ("偏高", "升高", "阳性", "弱阳性"):
        return "偏高"
    if explicit in ("偏低", "降低"):
        return "偏低"
    if explicit == "正常":
        return "正常"
    if explicit in ("异常", "阴性"):
        return "异常" if explicit == "异常" else "正常"

    try:
        number = float(value)
    except (TypeError, ValueError):
        return "异常"

    if low is not None and high is not None:
        if number < low:
            return "偏低"
        if number > high:
            return "偏高"
        return "正常"
    return "正常"


def _trim(num_text: str) -> str:
    try:
        number = float(num_text)
    except ValueError:
        return num_text
    return str(int(number)) if number == int(number) else str(number)


# ----------------------------------------------------------------------
# 元信息推断
# ----------------------------------------------------------------------
def guess_meta(text: str, filename: str) -> Dict[str, str]:
    """从文本与文件名推断报告名称与类型。"""
    haystack = f"{filename}\n{text[:3000]}"

    name = None
    for keyword in EXAM_KEYWORDS + LAB_KEYWORDS:  # 先匹配更具体的检查类
        if keyword in haystack:
            name = keyword
            break

    report_type = "laboratory"
    if name and name in EXAM_KEYWORDS:
        report_type = "examination"
    elif name is None:
        # 文件名兜底
        stem = Path(filename).stem.strip() or "上传报告"
        name = stem[:50]

    return {"report_name": name, "report_type": report_type}


def guess_department(text: str, report_name: str = "") -> Optional[str]:
    haystack = f"{report_name}\n{text[:2000]}"
    mapping = [
        ("血常规", "内科"), ("C反应蛋白", "内科"), ("肝功能", "消化科"),
        ("肾功能", "肾脏病中心"), ("肌酐", "肾脏病中心"), ("尿常规", "肾脏病中心"),
        ("血脂", "内科"), ("血糖", "内分泌科"), ("糖化血红蛋白", "内分泌科"),
        ("甲状腺", "内分泌科"), ("心肌酶", "内科"), ("心电图", "内科"),
        ("胸部CT", "呼吸科"), ("CT", "呼吸科"), ("胸片", "呼吸科"),
        ("妇科", "妇产科"), ("乳腺", "妇产科"),
    ]
    for keyword, department in mapping:
        if keyword in haystack:
            return department
    return None
